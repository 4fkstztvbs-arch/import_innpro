#!/usr/bin/env python3
"""Bose produkty: porovnanie s Heurekou a dopad zníženia ceny na 1. miesto na maržu.

Použitie: python3 scripts/bose-heureka-margin.py EXPORT.xml HEUREKA_REPORT.csv [--out=reports/bose-heureka-porovnanie-2026-10-15]
  * naše ceny a nákupky: export produktov z e-shopu (po importe), cena = akciová, ak je nižšia
  * konkurencia: Heureka sortiment report (stĺpce Najnižšia cena, PriceMin2..10), vlastná ponuka sa z poradia odstráni
  * 1. miesto: cena = najnižšia konkurenčná cena - 0,01 EUR (ak už sme najlacnejší, cena sa nemení)
"""
import csv, re, sys, os
from lxml import etree
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

VAT = 1.23
args = [a for a in sys.argv[1:] if not a.startswith("--")]
opts = dict(a[2:].split("=", 1) for a in sys.argv[1:] if a.startswith("--") and "=" in a)
export, report = args[0], args[1]
out = opts.get("out", "reports/bose-heureka-porovnanie-2026-10-15")
MIN_MARKUP = 0.05  # politika repozitára: minimálna prirážka 5 % nad nákupkou

def key(code):
    return re.sub(r"[\s_]+", " ", code or "").strip().upper()

def fnum(s):
    try: return float(str(s).replace(",", ".").strip())
    except ValueError: return None

heur = {}
with open(report, encoding="utf-8-sig", newline="") as f:
    for r in csv.DictReader(f):
        if "bose" in (r["Výrobca"] or "").lower() or "BASYS-B" in (r["Item ID"] or "").upper() and "bose" in r["Názov produktu"].lower():
            heur[key(r["Item ID"])] = r

rows, unmatched = [], []
for it in etree.parse(export).getroot().findall("SHOPITEM"):
    if it.find("VARIANTS") is not None: continue
    g = lambda p: (it.findtext(p) or "").strip()
    code = g("CODE")
    if g("AVAILABILITY_OUT_OF_STOCK") == "Vypredané" and g("VISIBILITY") == "detailOnly": continue
    price, act = fnum(g("PRICE_VAT")), fnum(g("ACTION_PRICE"))
    cur = act if act and act < price else price
    pur = fnum(g("PURCHASE_PRICE")) or 0
    if g("PURCHASE_PRICE_INCL_VAT") == "1": pur = pur / VAT
    h = heur.get(key(code))
    if not h:
        unmatched.append((code, g("NAME"))); continue
    mine = fnum(h["Vaša cena"])
    ranks = [fnum(h[c]) for c in ["Najnižšia cena"] + [f"PriceMin{i}" for i in range(2, 11)] if fnum(h[c]) is not None]
    comp = list(ranks)
    if mine in comp: comp.remove(mine)  # odstrániť vlastnú ponuku
    cmin = min(comp) if comp else None
    sellers = int(fnum(h["E-shopov predávajúcich produkt"]) or 0)
    pos_now = 1 + sum(1 for c in comp if c < cur - 0.004) if comp else 1
    if cmin is None:
        target = cur
    elif cur < cmin - 0.004:
        target = cur            # už sme najlacnejší
    else:
        target = round(cmin - 0.01, 2)
    # scenár B: cena sa zníži najviac na politikou povolenú prirážku (min. 5 % nad nákupkou), 1. miesto len ak je dosiahnuteľné
    floor_p = round(pur * (1 + MIN_MARKUP) * VAT + 0.004, 2) if pur else 0
    tgt_b = max(target, min(cur, floor_p)) if target < cur else cur
    m_now = cur / VAT - pur
    m_tgt = target / VAT - pur
    rows.append(dict(
        code=code, name=g("NAME"), ean=g("EAN"), seen_price=mine, cur=cur, pur=round(pur, 2),
        sellers=sellers, pos_report=int(fnum(h["Vaša pozícia podľa ceny"]) or 0), pos_now=pos_now,
        cmin=cmin, target=target, drop=round(cur - target, 2),
        m_now=round(m_now, 2), m_now_pct=round(m_now / (cur / VAT) * 100, 1), mk_now=round(m_now / pur * 100, 1) if pur else None,
        m_tgt=round(m_tgt, 2), m_tgt_pct=round(m_tgt / (target / VAT) * 100, 1), mk_tgt=round(m_tgt / pur * 100, 1) if pur else None,
        d_m=round(m_tgt - m_now, 2), d_pct=round(m_tgt / (target / VAT) * 100 - m_now / (cur / VAT) * 100, 1),
        below_floor=target / VAT < pur * (1 + MIN_MARKUP) - 0.005, below_cost=target / VAT < pur - 0.005,
        floor_p=floor_p, tgt_b=tgt_b, m_b=round(tgt_b / VAT - pur, 2), reach_b=(cmin is None) or tgt_b <= round(cmin - 0.01, 2) + 0.004 or cur < (cmin or 0) - 0.004,
        url=h["Heureka URL"]))
rows.sort(key=lambda r: (r["drop"] == 0, -r["drop"]))

hdr = ["Kód", "Produkt", "Predajcov na Heureke", "Naša cena dnes (s DPH)", "Nákupka (bez DPH)", "Marža dnes (€ bez DPH)", "Marža dnes (%)",
       "Pozícia dnes", "Najnižšia konkurenčná cena", "Cena pre 1. miesto", "Zníženie ceny (€)",
       "Marža po znížení (€ bez DPH)", "Marža po znížení (%)", "Pokles marže (€ / ks)", "Pokles marže (p. b.)", "Prirážka po znížení (%)", "Upozornenie", "Scenár B: cena s limitom 5 % prirážky", "Scenár B: marža (€ bez DPH)", "Scenár B: pokles marže (€ / ks)", "Scenár B: dosiahne 1. miesto", "Heureka"]
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Bose vs Heureka"
ws.append(hdr)
for r in rows:
    warn = "pod nákupkou" if r["below_cost"] else ("pod 5 % prirážkou" if r["below_floor"] else ("už 1. miesto" if r["drop"] == 0 else ""))
    ws.append([r["code"], r["name"], r["sellers"], r["cur"], r["pur"], r["m_now"], r["m_now_pct"], r["pos_now"], r["cmin"], r["target"], r["drop"],
               r["m_tgt"], r["m_tgt_pct"], r["d_m"], r["d_pct"], r["mk_tgt"], warn, r["tgt_b"], r["m_b"], round(r["m_b"] - r["m_now"], 2), "áno" if r["reach_b"] else "nie", r["url"]])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864"); c.alignment = Alignment(wrap_text=True, vertical="center")
for i, w in enumerate([24, 60, 11, 12, 12, 12, 10, 9, 14, 12, 11, 14, 11, 12, 11, 12, 18, 14, 14, 14, 12, 40], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions
red = PatternFill("solid", fgColor="F8CBAD")
for row in ws.iter_rows(min_row=2):
    if row[16].value in ("pod nákupkou", "pod 5 % prirážkou"):
        for c in row: c.fill = red
os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
wb.save(out + ".xlsx")
with open(out + ".csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f, delimiter=";"); w.writerow(hdr)
    for row in ws.iter_rows(min_row=2, values_only=True): w.writerow(row)

act = [r for r in rows if r["drop"] > 0]
tot_now = sum(r["m_now"] for r in rows); tot_tgt = sum(r["m_tgt"] for r in rows)
print(f"produktov: {len(rows)} (bez párovania: {len(unmatched)}), už na 1. mieste: {len(rows) - len(act)}, treba znížiť: {len(act)}")
print(f"súčet marže 1 ks každého produktu: dnes {tot_now:.2f} EUR -> {tot_tgt:.2f} EUR ({tot_tgt - tot_now:+.2f} EUR, {(tot_tgt / tot_now - 1) * 100:+.1f} %)")
print(f"priemerné zníženie ceny: {sum(r['drop'] for r in act) / max(1, len(act)):.2f} EUR; pod nákupkou: {sum(r['below_cost'] for r in act)}; pod 5 % prirážkou: {sum(r['below_floor'] and not r['below_cost'] for r in act)}")
for r in act[:40]:
    print(f"{r['code'][-16:]:16} {r['name'][:38]:38} {r['cur']:7.2f} -> {r['target']:7.2f}  marža {r['m_now']:6.2f} -> {r['m_tgt']:6.2f} ({r['d_m']:+.2f})  {'POD NÁKUPKOU' if r['below_cost'] else ('<5%' if r['below_floor'] else '')}")
b_tot = sum(r["m_b"] for r in rows)
print(f"SCENÁR B (limit 5 % prirážky): marža dnes {tot_now:.2f} -> {b_tot:.2f} EUR ({b_tot - tot_now:+.2f}); 1. miesto dosiahne {sum(1 for r in act if r['reach_b'])} z {len(act)} produktov")
print("nespárované:", unmatched[:10])
